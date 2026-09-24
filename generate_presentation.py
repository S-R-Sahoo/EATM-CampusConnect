import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_eatm_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # EATM Institutional Palette
    DARK_GREEN = RGBColor(11, 70, 39)     # #0B4627
    FOREST_GREEN = RGBColor(15, 81, 50)   # #0F5132
    EMERALD = RGBColor(16, 185, 129)      # #10B981
    MINT_BG = RGBColor(244, 253, 248)     # #F4FDF8
    CARD_BG = RGBColor(255, 255, 255)     # White
    CARD_BORDER = RGBColor(220, 240, 230) # Soft green border
    DARK_TEXT = RGBColor(30, 41, 59)      # #1E293B
    MUTED_TEXT = RGBColor(100, 116, 139)  # #64748B
    GOLD_ACCENT = RGBColor(217, 119, 6)   # #D97706
    LIGHT_GRAY_BG = RGBColor(248, 250, 252)

    blank_slide_layout = prs.slide_layouts[6]

    def add_background(slide, color=MINT_BG):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, subtitle_text="EATM CampusConnect • Next-Gen Academic & Social Ecosystem"):
        # Top Institutional Accent Bar
        top_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12)
        )
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = DARK_GREEN
        top_bar.line.fill.background()

        # Header Title Box
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        # Subtitle / Category
        p_sub = tf.paragraphs[0]
        p_sub.text = subtitle_text.upper()
        p_sub.font.size = Pt(10)
        p_sub.font.bold = True
        p_sub.font.color.rgb = EMERALD
        p_sub.font.name = "Arial"
        p_sub.space_after = Pt(4)

        # Title
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = DARK_GREEN
        p_title.font.name = "Georgia"

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
        )
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        return card

    def add_footer(slide, current_page=None, total_pages=14):
        # Footer Line & Text
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.4))
        tf = tb.text_frame
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        page_str = f"Slide {current_page} of {total_pages}" if current_page else ""
        p.text = f"Einstein Academy of Technology & Management (EATM)  |  Confidential & Institutional Project  |  {page_str}"
        p.font.size = Pt(9)
        p.font.color.rgb = MUTED_TEXT
        p.font.name = "Arial"

    # ==========================================================
    # SLIDE 1: Title Slide (Hero Dark Green)
    # ==========================================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    add_background(s1, DARK_GREEN)

    # Hero Accent Shape
    hero_accent = s1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(0.15), Inches(5.8)
    )
    hero_accent.fill.solid()
    hero_accent.fill.fore_color.rgb = EMERALD
    hero_accent.line.fill.background()

    # Title & Subtitle Box
    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(11.0), Inches(4.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "OFFICIAL INSTITUTIONAL PLATFORM SPECIFICATION"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = EMERALD
    p_badge.font.name = "Arial"
    p_badge.space_after = Pt(14)

    p_main = tf1.add_paragraph()
    p_main.text = "EATM CampusConnect"
    p_main.font.size = Pt(44)
    p_main.font.bold = True
    p_main.font.color.rgb = RGBColor(255, 255, 255)
    p_main.font.name = "Georgia"
    p_main.space_after = Pt(10)

    p_desc = tf1.add_paragraph()
    p_desc.text = "Next-Generation Campus Social & Academic Collaboration Ecosystem"
    p_desc.font.size = Pt(20)
    p_desc.font.color.rgb = RGBColor(220, 240, 230)
    p_desc.font.name = "Arial"
    p_desc.space_after = Pt(36)

    p_meta = tf1.add_paragraph()
    p_meta.text = "Institution: Einstein Academy of Technology & Management (EATM), Bhubaneswar\nArchitecture: React 18 • TypeScript • Tailwind CSS • Supabase Realtime • PostgreSQL 15\nPrepared For: Faculty, Academic Leadership, Students & Evaluation Committee"
    p_meta.font.size = Pt(11)
    p_meta.font.color.rgb = RGBColor(167, 243, 208)
    p_meta.font.name = "Arial"

    # ==========================================================
    # SLIDE 2: Executive Summary & Vision
    # ==========================================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    add_background(s2)
    add_header(s2, "Executive Summary & Platform Vision")
    add_footer(s2, 2)

    # 3 Strategic Pillars Cards
    pillars = [
        ("Institutional Unification", "Consolidates fragmented campus communication across students, faculty, and administration into a single secure, verified intranet portal.", "🏛️"),
        ("Social & Academic Synergy", "Combines rich peer discovery, project showcase feeds, and WhatsApp-grade messaging with verified study materials and assignment management.", "💡"),
        ("Engineering Excellence", "Engineered with client-first optimistic UI, zero-latency caching, sub-millisecond switching, Supabase Realtime, and robust Role-Based Security.", "⚡")
    ]
    for i, (title, desc, icon) in enumerate(pillars):
        c_left = Inches(0.8 + i * 3.95)
        add_card(s2, c_left, Inches(1.8), Inches(3.75), Inches(4.8))
        
        tb = s2.shapes.add_textbox(c_left + Inches(0.3), Inches(2.1), Inches(3.15), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_icon = tf.paragraphs[0]
        p_icon.text = icon
        p_icon.font.size = Pt(28)
        p_icon.space_after = Pt(12)

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(18)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_GREEN
        p_t.font.name = "Georgia"
        p_t.space_after = Pt(12)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = DARK_TEXT
        p_d.font.name = "Arial"

    # ==========================================================
    # SLIDE 3: The Problem Statement (Current Campus Gaps)
    # ==========================================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    add_background(s3)
    add_header(s3, "The Problem: Critical Campus Communication Gaps")
    add_footer(s3, 3)

    challenges = [
        ("Fragmented WhatsApp Groups", "Crucial college notices and assignment updates get lost inside unofficial, noisy chat groups with no centralized indexing or archive.", "⚠️"),
        ("Zero Cross-Department Discovery", "Engineering students have no structured directory to discover peers across branches (CSE, EE, Civil, Mech) for hackathons and projects.", "🔍"),
        ("Unverified Study Notes Sharing", "Lecture materials and question banks are scattered over third-party cloud drives without semester, branch, or faculty validation.", "📚"),
        ("No Role Separation & Governance", "Lack of institutional verification creates privacy concerns, impersonation risks, and chaotic student-teacher interactions.", "🔒")
    ]
    for i, (title, desc, icon) in enumerate(challenges):
        col = i % 2
        row = i // 2
        c_left = Inches(0.8 + col * 5.95)
        c_top = Inches(1.8 + row * 2.5)
        add_card(s3, c_left, c_top, Inches(5.75), Inches(2.25))

        tb = s3.shapes.add_textbox(c_left + Inches(0.3), c_top + Inches(0.25), Inches(5.15), Inches(1.75))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = f"{icon}  {title}"
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_GREEN
        p_t.font.name = "Georgia"
        p_t.space_after = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = DARK_TEXT
        p_d.font.name = "Arial"

    # ==========================================================
    # SLIDE 4: The Solution (EATM CampusConnect Ecosystem)
    # ==========================================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    add_background(s4)
    add_header(s4, "The Solution: A Unified Digital Campus Ecosystem")
    add_footer(s4, 4)

    # 4 Quadrants of the Solution
    sol_cards = [
        ("Institutional Social Feed", "Rich multimedia posts, project showcases, markdown code snippets, bookmarks, upvotes & verified comments.", "📰"),
        ("Ultra-Fast Realtime Chat", "WhatsApp-grade instant messaging, live waveform audio recording, media lightbox, typing indicators & presence.", "💬"),
        ("Smart Peer Discovery Directory", "Intelligent search & filtering by Branch, Semester, Technical Skills, and Hackathon team requirements.", "👥"),
        ("Academic & Resource Portal", "Faculty-approved study materials repository, assignment submission workflows, and campus event tracking.", "🎓")
    ]
    for i, (title, desc, icon) in enumerate(sol_cards):
        col = i % 2
        row = i // 2
        c_left = Inches(0.8 + col * 5.95)
        c_top = Inches(1.8 + row * 2.5)
        add_card(s4, c_left, c_top, Inches(5.75), Inches(2.25))

        tb = s4.shapes.add_textbox(c_left + Inches(0.3), c_top + Inches(0.25), Inches(5.15), Inches(1.75))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = f"{icon}  {title}"
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_GREEN
        p_t.font.name = "Georgia"
        p_t.space_after = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = DARK_TEXT
        p_d.font.name = "Arial"

    # ==========================================================
    # SLIDE 5: Stakeholder Profiles & Multi-Role Architecture
    # ==========================================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    add_background(s5)
    add_header(s5, "Multi-Role Architecture & Stakeholder Value")
    add_footer(s5, 5)

    roles = [
        ("Students", [
            "Discover tech partners & build hackathon teams",
            "Exchange notes & download verified lecture PDFs",
            "Showcase projects & receive peer endorsements",
            "Direct 1-on-1 and study group chat with classmates"
        ], "👨‍🎓"),
        ("Faculty & Mentors", [
            "Broadcast official departmental announcements",
            "Publish syllabi, lecture slides & lab manuals",
            "Track assignment submissions & review progress",
            "Offer direct guidance and mentorship to students"
        ], "👩‍🏫"),
        ("College Administration", [
            "Campus-wide broadcast alerts & event notifications",
            "Moderation tools ensuring safe institutional discourse",
            "Digital repository archiving campus academic assets",
            "Elimination of unmonitored communication silos"
        ], "🏛️")
    ]
    for i, (title, points, icon) in enumerate(roles):
        c_left = Inches(0.8 + i * 3.95)
        add_card(s5, c_left, Inches(1.8), Inches(3.75), Inches(4.8))

        tb = s5.shapes.add_textbox(c_left + Inches(0.25), Inches(2.0), Inches(3.25), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = f"{icon} {title}"
        p_t.font.size = Pt(18)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_GREEN
        p_t.font.name = "Georgia"
        p_t.space_after = Pt(12)

        for pt in points:
            p_bullet = tf.add_paragraph()
            p_bullet.text = f"•  {pt}"
            p_bullet.font.size = Pt(11)
            p_bullet.font.color.rgb = DARK_TEXT
            p_bullet.font.name = "Arial"
            p_bullet.space_after = Pt(8)

    # ==========================================================
    # SLIDE 6: Feature Spotlight: Campus Feed & Project Showcase
    # ==========================================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    add_background(s6)
    add_header(s6, "Feature Spotlight: Interactive Campus Feed")
    add_footer(s6, 6)

    # Left Column: Key Highlights
    add_card(s6, Inches(0.8), Inches(1.8), Inches(5.75), Inches(4.8))
    tb_feed_left = s6.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.15), Inches(4.2))
    tf_fl = tb_feed_left.text_frame
    tf_fl.word_wrap = True
    
    p = tf_fl.paragraphs[0]
    p.text = "Dynamic Community Stream"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    p.font.name = "Georgia"
    p.space_after = Pt(12)

    feed_bullets = [
        "Rich Post Creation: Supports rich text, embedded code blocks with syntax highlighting, external GitHub repo links, and multi-image uploads.",
        "Interactive Social Engagement: Instant like counts, bookmarks for quick revision, deep nested comments, and clean shareable URLs.",
        "Topic & Department Tagging: Categorized under #CSE, #Robotics, #Placements, #Hackathon, and #AcademicNotices.",
        "Optimistic UI Updates: Likes, comments, and saves react in 0ms on the client while syncing seamlessly with Supabase in the background."
    ]
    for b in feed_bullets:
        pb = tf_fl.add_paragraph()
        pb.text = f"✔  {b}"
        pb.font.size = Pt(11)
        pb.font.color.rgb = DARK_TEXT
        pb.font.name = "Arial"
        pb.space_after = Pt(8)

    # Right Column: Project Showcase Cards
    add_card(s6, Inches(6.75), Inches(1.8), Inches(5.75), Inches(4.8), bg_color=LIGHT_GRAY_BG)
    tb_feed_right = s6.shapes.add_textbox(Inches(7.05), Inches(2.1), Inches(5.15), Inches(4.2))
    tf_fr = tb_feed_right.text_frame
    tf_fr.word_wrap = True

    p2 = tf_fr.paragraphs[0]
    p2.text = "Student Innovation & Recognition"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = FOREST_GREEN
    p2.font.name = "Georgia"
    p2.space_after = Pt(12)

    showcase_bullets = [
        "Project Showcases: Students can pin capstone projects, research papers, and software tools to their verified campus profile.",
        "Peer Endorsements: Classmates and professors can endorse technical proficiencies (Python, IoT, Web3, VLSI, AI/ML).",
        "Campus Recognition: Outstanding projects are surfaced on trending feeds, attracting recruiters and faculty project grants.",
        "Portfolio Generation: Automatically compiles a verifiable academic portfolio showcasing real campus contributions."
    ]
    for b in showcase_bullets:
        pb = tf_fr.add_paragraph()
        pb.text = f"★  {b}"
        pb.font.size = Pt(11)
        pb.font.color.rgb = DARK_TEXT
        pb.font.name = "Arial"
        pb.space_after = Pt(8)

    # ==========================================================
    # SLIDE 7: Feature Spotlight: WhatsApp-Grade Realtime Messaging
    # ==========================================================
    s7 = prs.slides.add_slide(blank_slide_layout)
    add_background(s7)
    add_header(s7, "Feature Spotlight: Realtime Campus Messaging")
    add_footer(s7, 7)

    chat_features = [
        ("Instant 0ms Messaging", "Ultra-low latency chat powered by Supabase WebSocket subscriptions and client-side deduplication.", "⚡"),
        ("Live Audio Voice Notes", "Record and preview voice notes with real-time waveform bars, pause/resume, and instant playback player.", "🎙️"),
        ("Media & Lightbox Gallery", "High-res photo, video, and PDF document sharing with full-screen lightbox modal viewer.", "🖼️"),
        ("Presence & Typing Indicators", "Real-time 'Online' / 'Last seen' timestamps and animated 3-dot typing alerts across active peers.", "🟢"),
        ("WhatsApp Badges & Auto-Clear", "Green numeric unread counter pills on conversations with instant automatic clearing upon viewing.", "📬"),
        ("Locked Mobile Fixed Header", "Fixed top header pinned at safe-area top on smartphones, remaining rock-solid when virtual keyboard opens.", "📱")
    ]
    for i, (title, desc, icon) in enumerate(chat_features):
        col = i % 3
        row = i // 3
        c_left = Inches(0.8 + col * 3.95)
        c_top = Inches(1.8 + row * 2.5)
        add_card(s7, c_left, c_top, Inches(3.75), Inches(2.25))

        tb = s7.shapes.add_textbox(c_left + Inches(0.25), c_top + Inches(0.2), Inches(3.25), Inches(1.85))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = f"{icon}  {title}"
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_GREEN
        p_t.font.name = "Georgia"
        p_t.space_after = Pt(6)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = DARK_TEXT
        p_d.font.name = "Arial"

    # ==========================================================
    # SLIDE 8: Feature Spotlight: Peer Discovery & Connections
    # ==========================================================
    s8 = prs.slides.add_slide(blank_slide_layout)
    add_background(s8)
    add_header(s8, "Feature Spotlight: Peer Discovery & Connections")
    add_footer(s8, 8)

    add_card(s8, Inches(0.8), Inches(1.8), Inches(5.75), Inches(4.8))
    tb_disc_left = s8.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.15), Inches(4.2))
    tf_dl = tb_disc_left.text_frame
    tf_dl.word_wrap = True

    p = tf_dl.paragraphs[0]
    p.text = "Intelligent Campus Directory"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    p.font.name = "Georgia"
    p.space_after = Pt(12)

    disc_bullets = [
        "Multi-Facet Search: Filter students and faculty by Branch (CSE, ECE, EE, MECH, CIVIL), Graduation Year, and Academic Role.",
        "Skill-Based Matching: Search for peers proficient in specific frameworks (React, Flutter, Solidity, Embedded C, TensorFlow).",
        "Connection Request Lifecycle: Send, accept, or decline connection requests with real-time status updates.",
        "Direct Message Integration: Accepted connections automatically link into the chat directory for immediate messaging."
    ]
    for b in disc_bullets:
        pb = tf_dl.add_paragraph()
        pb.text = f"✔  {b}"
        pb.font.size = Pt(11)
        pb.font.color.rgb = DARK_TEXT
        pb.font.name = "Arial"
        pb.space_after = Pt(8)

    add_card(s8, Inches(6.75), Inches(1.8), Inches(5.75), Inches(4.8), bg_color=LIGHT_GRAY_BG)
    tb_disc_right = s8.shapes.add_textbox(Inches(7.05), Inches(2.1), Inches(5.15), Inches(4.2))
    tf_dr = tb_disc_right.text_frame
    tf_dr.word_wrap = True

    p2 = tf_dr.paragraphs[0]
    p2.text = "Academic & Hackathon Networking"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = FOREST_GREEN
    p2.font.name = "Georgia"
    p2.space_after = Pt(12)

    net_bullets = [
        "Hackathon Team Builder: Post vacancies for UI designers, backend engineers, or hardware experts for upcoming competitions.",
        "Cross-Batch Mentorship: Junior students can connect with final-year seniors for placement prep, GATE coaching, and guidance.",
        "Study Group Creation: Form direct or cohort-based study groups to prepare for semester exams and lab evaluations.",
        "Privacy Safeguards: Students retain full control over connection requests with instant block/remove options."
    ]
    for b in net_bullets:
        pb = tf_dr.add_paragraph()
        pb.text = f"★  {b}"
        pb.font.size = Pt(11)
        pb.font.color.rgb = DARK_TEXT
        pb.font.name = "Arial"
        pb.space_after = Pt(8)

    # ==========================================================
    # SLIDE 9: Feature Spotlight: Academic Hub & Study Repository
    # ==========================================================
    s9 = prs.slides.add_slide(blank_slide_layout)
    add_background(s9)
    add_header(s9, "Feature Spotlight: Academic Hub & Materials")
    add_footer(s9, 9)

    acad_cards = [
        ("Verified Study Materials", "Centralized repository of lecture notes, question banks, lab sheets, and reference textbooks categorized by Department and Semester.", "📖"),
        ("Faculty Assignment Desk", "Teachers upload assignment briefs with deadlines; students upload submissions with automated timestamping.", "📝"),
        ("Institutional Moderation", "All public resources are peer-reviewed or faculty-approved to maintain high academic rigor and prevent outdated material.", "🛡️"),
        ("Fast One-Click Downloads", "Cloud-backed PDF, document, and media downloads with file size indicators and offline local caching.", "⬇️")
    ]
    for i, (title, desc, icon) in enumerate(acad_cards):
        col = i % 2
        row = i // 2
        c_left = Inches(0.8 + col * 5.95)
        c_top = Inches(1.8 + row * 2.5)
        add_card(s9, c_left, c_top, Inches(5.75), Inches(2.25))

        tb = s9.shapes.add_textbox(c_left + Inches(0.3), c_top + Inches(0.25), Inches(5.15), Inches(1.75))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = f"{icon}  {title}"
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_GREEN
        p_t.font.name = "Georgia"
        p_t.space_after = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = DARK_TEXT
        p_d.font.name = "Arial"

    # ==========================================================
    # SLIDE 10: Events, Hackathons & Communities
    # ==========================================================
    s10 = prs.slides.add_slide(blank_slide_layout)
    add_background(s10)
    add_header(s10, "Feature Spotlight: Events & Student Communities")
    add_footer(s10, 10)

    comm_cards = [
        ("Campus Events & Tech Fests", "Browse upcoming college symposiums, coding competitions, sports meets, and cultural festivals with one-click RSVP.", "🎪"),
        ("Interest-Based Communities", "Join or create departmental societies, coding clubs, robotics forums, and competitive programming circles.", "🚀"),
        ("Internships & Job Board", "Curated placement notices, campus drive schedules, alumni job referrals, and summer training alerts.", "💼"),
        ("Live Alert Notifications", "Synchronous high-priority bell alerts and unread counters for important administrative and club announcements.", "🔔")
    ]
    for i, (title, desc, icon) in enumerate(comm_cards):
        col = i % 2
        row = i // 2
        c_left = Inches(0.8 + col * 5.95)
        c_top = Inches(1.8 + row * 2.5)
        add_card(s10, c_left, c_top, Inches(5.75), Inches(2.25))

        tb = s10.shapes.add_textbox(c_left + Inches(0.3), c_top + Inches(0.25), Inches(5.15), Inches(1.75))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = f"{icon}  {title}"
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_GREEN
        p_t.font.name = "Georgia"
        p_t.space_after = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = DARK_TEXT
        p_d.font.name = "Arial"

    # ==========================================================
    # SLIDE 11: System Architecture & Data Flow
    # ==========================================================
    s11 = prs.slides.add_slide(blank_slide_layout)
    add_background(s11)
    add_header(s11, "System Architecture & Engineering Design")
    add_footer(s11, 11)

    arch_layers = [
        ("Presentation Layer", "• React 18 SPA with TypeScript\n• Tailwind CSS & Responsive Glassmorphism\n• Lucide Icons & Custom Audio Waveform\n• Safe-Area Viewport Pinned Controls", "💻"),
        ("State & Cache Layer", "• Auth & Theme Context Providers\n• In-Memory Fast Cache (30s TTL)\n• LocalStorage Synchronous Optimistic Fallback\n• Client-Side Message Deduplication Engine", "⚡"),
        ("Backend & Realtime", "• Supabase Managed PostgreSQL 15\n• Realtime WebSocket Replication Bus\n• Row Level Security (RLS) Policy Gates\n• Supabase Storage (Media, Docs, Audio)", "🗄️")
    ]
    for i, (title, bullets, icon) in enumerate(arch_layers):
        c_left = Inches(0.8 + i * 3.95)
        add_card(s11, c_left, Inches(1.8), Inches(3.75), Inches(4.8))

        tb = s11.shapes.add_textbox(c_left + Inches(0.25), Inches(2.0), Inches(3.25), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = f"{icon} {title}"
        p_t.font.size = Pt(17)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_GREEN
        p_t.font.name = "Georgia"
        p_t.space_after = Pt(12)

        for line in bullets.split("\n"):
            p_b = tf.add_paragraph()
            p_b.text = line
            p_b.font.size = Pt(11)
            p_b.font.color.rgb = DARK_TEXT
            p_b.font.name = "Arial"
            p_b.space_after = Pt(6)

    # ==========================================================
    # SLIDE 12: Technology Stack Summary
    # ==========================================================
    s12 = prs.slides.add_slide(blank_slide_layout)
    add_background(s12)
    add_header(s12, "Technology Stack & Engineering Highlights")
    add_footer(s12, 12)

    tech_categories = [
        ("Frontend Technologies", "React 18, TypeScript, Vite 6, Tailwind CSS 3.4, Lucide React, HTML5 Audio API, MediaRecorder API.", "⚛️"),
        ("Database & Cloud Services", "PostgreSQL 15, Supabase Database, Supabase Auth, Supabase Storage, Realtime WebSocket Channels.", "🐘"),
        ("Performance & Caching", "Optimistic Client Updates, 0ms Cache First Load, Message Hash Deduplication, 100dvh Dynamic Viewports.", "🚀"),
        ("Code Quality & Tooling", "Strict TypeScript Type-Safety, ESLint, Git Version Control, Vite Rollup Production Bundler.", "🛠️")
    ]
    for i, (title, desc, icon) in enumerate(tech_categories):
        col = i % 2
        row = i // 2
        c_left = Inches(0.8 + col * 5.95)
        c_top = Inches(1.8 + row * 2.5)
        add_card(s12, c_left, c_top, Inches(5.75), Inches(2.25))

        tb = s12.shapes.add_textbox(c_left + Inches(0.3), c_top + Inches(0.25), Inches(5.15), Inches(1.75))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = f"{icon}  {title}"
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_GREEN
        p_t.font.name = "Georgia"
        p_t.space_after = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = DARK_TEXT
        p_d.font.name = "Arial"

    # ==========================================================
    # SLIDE 13: Security, Privacy & Compliance
    # ==========================================================
    s13 = prs.slides.add_slide(blank_slide_layout)
    add_background(s13)
    add_header(s13, "Security, Privacy & Data Governance")
    add_footer(s13, 13)

    security_cards = [
        ("Row Level Security (RLS)", "PostgreSQL RLS ensures students can only read and write data that belongs to their verified identity, preventing unauthorized database queries.", "🔒"),
        ("Role-Based Access Control", "Rigid role permissions distinguish between Students, Faculty, and Admin — preventing privilege escalation across sections.", "🛡️"),
        ("Encrypted File Storage", "All uploaded assignment documents, project media, and voice notes are stored in secure Supabase Storage buckets with MIME validation.", "📁"),
        ("Privacy & Moderation", "Students can mute conversations, delete messages for themselves/everyone, and report inappropriate content for administrative review.", "👁️")
    ]
    for i, (title, desc, icon) in enumerate(security_cards):
        col = i % 2
        row = i // 2
        c_left = Inches(0.8 + col * 5.95)
        c_top = Inches(1.8 + row * 2.5)
        add_card(s13, c_left, c_top, Inches(5.75), Inches(2.25))

        tb = s13.shapes.add_textbox(c_left + Inches(0.3), c_top + Inches(0.25), Inches(5.15), Inches(1.75))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = f"{icon}  {title}"
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_GREEN
        p_t.font.name = "Georgia"
        p_t.space_after = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = DARK_TEXT
        p_d.font.name = "Arial"

    # ==========================================================
    # SLIDE 14: Future Roadmap & Conclusion (Hero Dark Green)
    # ==========================================================
    s14 = prs.slides.add_slide(blank_slide_layout)
    add_background(s14, DARK_GREEN)

    # Accent
    hero_accent2 = s14.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(0.15), Inches(5.8)
    )
    hero_accent2.fill.solid()
    hero_accent2.fill.fore_color.rgb = EMERALD
    hero_accent2.line.fill.background()

    tb14 = s14.shapes.add_textbox(Inches(1.2), Inches(1.0), Inches(11.0), Inches(5.5))
    tf14 = tb14.text_frame
    tf14.word_wrap = True

    p_c1 = tf14.paragraphs[0]
    p_c1.text = "STRATEGIC ROADMAP & CONCLUSION"
    p_c1.font.size = Pt(12)
    p_c1.font.bold = True
    p_c1.font.color.rgb = EMERALD
    p_c1.font.name = "Arial"
    p_c1.space_after = Pt(10)

    p_c2 = tf14.add_paragraph()
    p_c2.text = "Empowering EATM's Digital Future"
    p_c2.font.size = Pt(32)
    p_c2.font.bold = True
    p_c2.font.color.rgb = RGBColor(255, 255, 255)
    p_c2.font.name = "Georgia"
    p_c2.space_after = Pt(16)

    future_points = [
        "Phase 1 (Completed): Social Feed, Realtime Chat, Voice Notes, Peer Directory, Study Materials, and Multi-Role Dashboard.",
        "Phase 2 (Upcoming): AI-Powered Academic Tutor (Gemini API) for instant doubt solving and automated paper summarization.",
        "Phase 3 (Upcoming): College ERP Integration for automated semester attendance, timetable syncing, and grade book viewing.",
        "Phase 4 (Upcoming): Alumni Mentorship & Direct Industry Placement Pipeline with verified referral badges."
    ]
    for pt in future_points:
        p_pt = tf14.add_paragraph()
        p_pt.text = f"★  {pt}"
        p_pt.font.size = Pt(13)
        p_pt.font.color.rgb = RGBColor(220, 240, 230)
        p_pt.font.name = "Arial"
        p_pt.space_after = Pt(10)

    p_end = tf14.add_paragraph()
    p_end.text = "\nThank You!  |  EATM CampusConnect: Connecting Minds, Empowering Futures."
    p_end.font.size = Pt(14)
    p_end.font.bold = True
    p_end.font.color.rgb = RGBColor(167, 243, 208)
    p_end.font.name = "Georgia"

    # Save presentation
    output_path = os.path.join(os.getcwd(), "EATM_CampusConnect_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    build_eatm_presentation()
